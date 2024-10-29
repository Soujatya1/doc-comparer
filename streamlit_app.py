import streamlit as st
from langchain.document_loaders import PyPDFLoader
from langchain.embeddings import HuggingFaceEmbeddings
from langchain.vectorstores import FAISS
from langchain.chains import RetrievalQA
from langchain_groq import ChatGroq
import os
from langchain.chains.question_answering import load_qa_chain
from docx import Document as DocxDocument
from langchain.schema import Document as LangChainDocument

# Initialize the Streamlit app
st.title("Document Comparer!")
llm = ChatGroq(groq_api_key="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", model_name="Llama3-8b-8192")


def initialize_vectorstore(documents):
    # Initialize embeddings model
    embeddings = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")
    # Create FAISS vector store from documents
    vectorstore = FAISS.from_documents(documents, embeddings)
    return vectorstore

def load_docx(file_path):
    doc = DocxDocument(file_path)
    text = []
    for paragraph in doc.paragraphs:
        text.append(paragraph.text)
    return "\n".join(text)

uploaded_files = st.file_uploader("Upload PDF documents", accept_multiple_files=True, type=["text", "docx", "pdf"])

documents = []
if uploaded_files:
    for uploaded_file in uploaded_files:
        file_path = f"temp_{uploaded_file.name}"
        with open(file_path, "wb") as f:
            f.write(uploaded_file.read())

        # Check file type and load documents accordingly
        if uploaded_file.type == "application/pdf":
            loader = PyPDFLoader(file_path)
            pdf_documents = loader.load()
            documents.extend(pdf_documents)  # Extend with PDF documents

        elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc_text = load_docx(file_path)
            # Create a LangChain document object for the loaded text
            documents.append(LangChainDocument(page_content=doc_text, metadata={"source": file_path}))

    # Initialize FAISS vector store with all documents (no caching)
    vectorstore = initialize_vectorstore(documents)
    st.success("Documents uploaded and embedded successfully!")

    # Set up RetrievalQA chain with LangChain
    retriever = vectorstore.as_retriever()
    
    # Load a basic QA chain for combining documents
    combine_chain = load_qa_chain(llm, chain_type="stuff")

    qa_chain = RetrievalQA(combine_documents_chain=combine_chain, retriever=retriever)

    # Define comparison function
    def compare_documents(query):
        response = qa_chain({"query": query})
        return response["result"]

    # Interface to compare documents
    st.subheader("Compare Documents")
    query = st.text_input("Enter your comparison query:")

    if st.button("Compare"):
        if not query:
            st.warning("Please enter a query to compare.")
        else:
            result = compare_documents(query)
            st.write("Comparison Result:")
            st.write(result)
else:
    st.warning("Please upload documents to enable comparison.")

# Clean up temporary files
for uploaded_file in uploaded_files:
    os.remove(f"temp_{uploaded_file.name}")
